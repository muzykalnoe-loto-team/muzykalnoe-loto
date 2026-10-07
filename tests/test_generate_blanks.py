import pytest
import openpyxl
from scripts.generate_blanks import generate_blanks
from scripts import paths
from scripts.song_bank import load_song_bank

TEMPLATE = paths.BLANKS_TEMPLATE
BANK = paths.SONG_BANK


@pytest.mark.integration
def test_generate_blanks_fills_individual_sheet(tmp_path):
    bank = load_song_bank(BANK)
    selected = [item["raw"] for item in bank[:90]]

    output = tmp_path / "blanks_output.xlsm"

    generate_blanks(TEMPLATE, str(output), selected, n_individual=3, n_team=2)

    assert output.exists()

    wb = openpyxl.load_workbook(output, data_only=True)
    ws = wb["Индивидуальные бланки"]
    # Первый бланк начинается в A1 с номером бланка в I1
    assert ws["A1"].value == "Музыкальное лото\nvk.com/dvizh_nv"
    assert ws["I1"].value == 1
    # Хотя бы одна ячейка сетки заполнена реальной песней, а не пустая
    assert ws["A3"].value not in (None, "")


@pytest.mark.integration
def test_generate_blanks_rejects_wrong_song_count(tmp_path):
    output = tmp_path / "blanks_output.xlsm"

    with pytest.raises(ValueError, match="90"):
        generate_blanks(TEMPLATE, str(output), ["Артист - Песня"], n_individual=1, n_team=1)
