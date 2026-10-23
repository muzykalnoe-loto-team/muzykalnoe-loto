import json
import os
import shutil

import openpyxl
import pytest
from pptx import Presentation

from scripts import paths
from scripts.export_selection import export_selection
from scripts.run_pipeline import run_pipeline

# 60 видео и 30 старых слайдов (бочонки 3, 6 … 90)
SELECTION = "tests/data/список — смешанный.json"


@pytest.mark.integration
def test_run_pipeline_produces_both_files(tmp_path):
    selected_path = tmp_path / "selected_songs.json"
    shutil.copy(SELECTION, selected_path)

    result = run_pipeline(str(selected_path), n_individual=3, n_team=2, output_dir=str(tmp_path))

    blanks = tmp_path / "Бланки.xlsm"
    presentation = tmp_path / "Презентация.pptx"
    assert result == {"blanks_path": str(blanks), "presentation_path": str(presentation), "warning": "",
                      "nominations": None}

    wb = openpyxl.load_workbook(blanks, data_only=True)
    with open(SELECTION, encoding="utf-8") as f:
        barrels = json.load(f)["barrels"]
    # в шаблон бланков ушли все 90 песен по порядку бочонков
    assert wb["Список"]["B2"].value == barrels[0]["song"]
    assert wb["Список"]["E1"].value == 3 and wb["Список"]["E2"].value == 2

    prs = Presentation(str(presentation))
    assert len(prs.slides) == 118
    # на слоте бочонка 89 (слайд 104) — видео нужной песни
    names = [s.name for s in prs.slides[103].shapes]
    assert any(n.startswith(barrels[88]["song"].split(" - ")[0]) for n in names)


def test_run_pipeline_rejects_bad_selection_before_office(tmp_path):
    selected_path = tmp_path / "selected_songs.json"
    with open(selected_path, "w", encoding="utf-8") as f:
        json.dump({"barrels": [{"number": 1, "song": "Несуществующий Артист - Песня"}]}, f, ensure_ascii=False)

    with pytest.raises(ValueError, match="90"):
        run_pipeline(str(selected_path), n_individual=1, n_team=1, output_dir=str(tmp_path))

    assert os.listdir(tmp_path) == ["selected_songs.json"]


def _common_bank():
    if not os.path.exists(paths.COMMON_BANK_JSON):
        pytest.skip("общий банк ещё не собран (Task 4)")
    from scripts.common_bank import load_common_bank
    return load_common_bank()


def test_all_video_list_is_rejected_before_office(tmp_path):
    selected_path = tmp_path / "selected_songs.json"
    shutil.copy("tests/data/список — только видео.json", selected_path)  # старый пример: 90 видео
    with pytest.raises(ValueError, match="Бочонок 3"):
        run_pipeline(str(selected_path), n_individual=1, n_team=1, output_dir=str(tmp_path))
    assert os.listdir(tmp_path) == ["selected_songs.json"]


def test_run_pipeline_rejects_old_slide_on_wrong_barrel(tmp_path):
    bank = _common_bank()
    old = [e["raw"] for e in bank if e["old_file"]][:30]
    video = [e["raw"] for e in bank if e["video_file"] and e["raw"] not in old][:60]
    songs, kinds = [], []
    for n in range(1, 91):
        is_old = n % 3 == 0
        songs.append(old.pop() if is_old else video.pop())
        kinds.append("old" if is_old else "video")
    kinds[3], kinds[2] = "old", "video"  # бочонок 4 — старый слайд, бочонок 3 — видео: ошибка
    selected_path = tmp_path / "selected_songs.json"
    export_selection(songs, str(selected_path), kinds=kinds)

    with pytest.raises(ValueError, match="Бочонок 3"):
        run_pipeline(str(selected_path), n_individual=1, n_team=1, output_dir=str(tmp_path))
    assert sorted(os.listdir(tmp_path)) == ["selected_songs.json"]


def test_earlier_event_list_is_found_in_common_bank():
    # список «Готического бала» собран до общего банка: часть песен в нём написана иначе (ё/е, регистр)
    from scripts.barrel_kinds import bank_index, check_kinds, resolve_files
    from scripts.validate_selection import validate_selection
    bank = _common_bank()
    with open("tests/data/Готический бал — список.json", encoding="utf-8") as f:
        barrels = json.load(f)["barrels"]

    assert validate_selection([b["song"] for b in barrels], bank)["errors"] == []
    index = bank_index(bank)
    # он из одних видео — такой список сборка теперь не принимает: на бочонках 3, 6 … 90 нужны старые слайды
    errors = check_kinds(barrels, index)
    assert len(errors) == 30 and all("старый слайд" in e for e in errors)
    for b in resolve_files(barrels, index):
        assert os.path.exists(os.path.join(paths.SLIDES_ARCHIVE, b["file"]))


def test_size_warning_when_presentation_is_too_big(tmp_path):
    from scripts.run_pipeline import size_warning
    big, small = tmp_path / "big.pptx", tmp_path / "small.pptx"
    big.write_bytes(b"0" * 2_000_000)
    small.write_bytes(b"0" * 500_000)

    assert "больше 1 МБ" in size_warning(str(big), limit_mb=1)  # лимит Яндекс Диска — 995 МБ
    assert size_warning(str(small), limit_mb=1) == ""


def test_themed_without_picture_is_rejected_before_office(tmp_path):
    selected = tmp_path / "selected_songs.json"
    shutil.copy(SELECTION, selected)
    with pytest.raises(ValueError, match="картинк"):
        run_pipeline(str(selected), n_individual=1, n_team=1, output_dir=str(tmp_path),
                     nominations={"themed": True, "choices": {}, "theme": None})
    assert os.listdir(tmp_path) == ["selected_songs.json"]


def test_unknown_nomination_is_rejected_before_office(tmp_path):
    selected = tmp_path / "selected_songs.json"
    shutil.copy(SELECTION, selected)
    with pytest.raises(ValueError, match="раунда"):
        run_pipeline(str(selected), n_individual=1, n_team=1, output_dir=str(tmp_path),
                     nominations={"themed": False, "choices": {"2": "sing"}, "theme": None})
    assert os.listdir(tmp_path) == ["selected_songs.json"]


@pytest.mark.integration
def test_pipeline_applies_nominations_and_records_history(tmp_path, monkeypatch):
    from scripts.nominations import load_history
    monkeypatch.setattr(paths, "NOMINATION_HISTORY", str(tmp_path / "h.json"))
    selected = tmp_path / "selected_songs.json"
    shutil.copy(SELECTION, selected)  # обычная игра, «Авто»; выбор из брифа приходит строками-ключами
    result = run_pipeline(str(selected), n_individual=1, n_team=1, output_dir=str(tmp_path),
                          event="Проба", event_date="2026-10-19",
                          nominations={"themed": False, "choices": {"3": "nobrakes"}, "theme": None})
    assert result["nominations"] == {1: "sing", 2: "goddess", 3: "nobrakes"}
    assert load_history(str(tmp_path / "h.json"))[-1]["event"] == "Проба"
    ids = [s.slide_id for s in Presentation(result["presentation_path"]).slides]
    assert len(ids) == 118 - 4 and {375, 259, 378} <= set(ids)


@pytest.mark.integration
def test_locked_archive_does_not_break_build(tmp_path, monkeypatch):
    # архив — копия шаблона во временной папке, открытая в PowerPoint на время сборки (как у оператора)
    import win32com.client as win32

    from scripts.theme_art import render_scene
    from scripts.theme_scenes import THEMES
    archive = tmp_path / "Номинация.pptx"
    shutil.copy(paths.PRESENTATION_TEMPLATE, archive)
    monkeypatch.setattr(paths, "NOMINATION_ARCHIVE", str(archive))
    monkeypatch.setattr(paths, "NOMINATION_HISTORY", str(tmp_path / "h.json"))
    art = render_scene(THEMES["gatsby"]["scene"], str(tmp_path / "g.png"), seed="gatsby")
    app = win32.Dispatch("PowerPoint.Application")
    held = app.Presentations.Open(str(archive), WithWindow=False)
    try:
        selected = tmp_path / "selected_songs.json"
        shutil.copy(SELECTION, selected)
        (tmp_path / "игра").mkdir()
        result = run_pipeline(str(selected), n_individual=1, n_team=1, output_dir=str(tmp_path / "игра"),
                              event="Гэтсби-бал", event_date="2026-11-01",
                              nominations={"themed": True, "choices": {}, "theme": {
                                  "title": THEMES["gatsby"]["title"], "glow": THEMES["gatsby"]["glow"],
                                  "art_png": art, "name": "Гэтсби"}})
    finally:
        held.Close()
        if app.Presentations.Count == 0:
            app.Quit()
    assert os.path.exists(result["presentation_path"]) and "архив" in result["warning"].lower()
    assert result["nominations"][1] == "theme"


def test_stop_list_song_is_rejected_before_office(tmp_path, monkeypatch):
    from openpyxl import Workbook
    with open(SELECTION, encoding="utf-8") as f:
        artist = json.load(f)["barrels"][0]["song"].split(" - ")[0]
    wb = Workbook()
    wb.active.append(["Исполнитель", "Причина"])
    wb.active.append([artist, "проверка"])
    wb.save(tmp_path / "стоп.xlsx")
    monkeypatch.setattr(paths, "STOP_LIST", str(tmp_path / "стоп.xlsx"))
    selected = tmp_path / "selected_songs.json"
    shutil.copy(SELECTION, selected)
    with pytest.raises(ValueError, match="стоп-лист"):
        run_pipeline(str(selected), n_individual=1, n_team=1, output_dir=str(tmp_path))
    assert sorted(os.listdir(tmp_path)) == ["selected_songs.json", "стоп.xlsx"]
