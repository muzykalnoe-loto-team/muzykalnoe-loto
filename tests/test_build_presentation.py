import json

import pytest
from pptx import Presentation

from scripts import paths
from scripts.build_presentation import assemble_presentation

TEMPLATE = paths.PRESENTATION_TEMPLATE
ARCHIVE = paths.SLIDES_ARCHIVE
SELECTION = f"{paths.EVENTS}/пример/selected_songs.json"


def _shape_names(prs, slide_number_1based):
    return [s.name for s in prs.slides[slide_number_1based - 1].shapes]


@pytest.mark.integration
def test_assemble_presentation_fills_requested_slots(tmp_path):
    with open(SELECTION, encoding="utf-8") as f:
        all_barrels = json.load(f)["barrels"]
    barrels = [all_barrels[0], all_barrels[1], all_barrels[89]]  # бочонки 1, 2, 90

    output = tmp_path / "presentation_output.pptx"

    assemble_presentation(TEMPLATE, ARCHIVE, barrels, str(output))

    prs = Presentation(str(output))
    # Общее число слайдов не изменилось — мы не удаляли и не добавляли слайды
    assert len(prs.slides) == 118
    # на слотах бочонков 1, 2, 90 (слайды 16, 17, 105) появилось видео песни
    for slide_number, barrel in [(16, barrels[0]), (17, barrels[1]), (105, barrels[2])]:
        assert barrel["song"] in _shape_names(prs, slide_number)
    # незаполненный слот (бочонок 3) остался как в шаблоне
    assert _shape_names(prs, 18) == _shape_names(Presentation(TEMPLATE), 18)


@pytest.mark.integration
def test_assemble_presentation_fails_fast_on_missing_archive_file(tmp_path):
    barrels = [{"number": 1, "song": "Несуществующий Артист - Несуществующая Песня"}]
    output = tmp_path / "presentation_output.pptx"

    with pytest.raises(FileNotFoundError, match="Несуществующий Артист"):
        assemble_presentation(TEMPLATE, ARCHIVE, barrels, str(output))

    assert not output.exists()


@pytest.mark.integration
def test_assemble_mixed_game_puts_old_slide_on_third_barrel(tmp_path):
    from pptx.enum.shapes import PP_PLACEHOLDER

    from scripts.barrel_kinds import bank_index, resolve_files
    from scripts.common_bank import load_common_bank
    index = bank_index(load_common_bank())
    video = next(e["raw"] for e in index.values() if e["video_file"])
    old = next(e["raw"] for e in index.values() if e["old_file"] and e["old_photo_part"] and e["raw"] != video)
    barrels = resolve_files([{"number": 1, "song": video, "kind": "video"},
                             {"number": 3, "song": old, "kind": "old"}], index)
    output = tmp_path / "presentation_output.pptx"

    assemble_presentation(TEMPLATE, ARCHIVE, barrels, str(output))

    slide = Presentation(str(output)).slides[17]  # бочонок 3
    assert any("audioFile" in s._element.xml for s in slide.shapes)  # звук старого слайда (python-pptx видит его картинкой)
    photo = next(p for p in slide.placeholders if p.placeholder_format.type == PP_PLACEHOLDER.PICTURE)
    assert photo.__class__.__name__ == "PlaceholderPicture"
    assert video in _shape_names(Presentation(str(output)), 16)
