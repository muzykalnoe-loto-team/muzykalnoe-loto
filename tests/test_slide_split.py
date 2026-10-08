import io
import os
import re
import zipfile

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Pt

from scripts import paths
from scripts.slide_split import extract_slide


def _png(color):
    buf = io.BytesIO()
    Image.new("RGB", (30, 20), color).save(buf, "PNG")
    buf.seek(0)
    return buf


@pytest.fixture
def three_slides(tmp_path):
    prs = Presentation()
    colors = [(255, 0, 0), (0, 255, 0), (0, 0, 255)]
    for i, color in enumerate(colors, 1):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.shapes.add_textbox(Pt(10), Pt(10), Pt(200), Pt(30)).text_frame.text = f"Слайд {i}"
        s.shapes.add_picture(_png(color), Pt(10), Pt(60), Pt(100), Pt(70))
    # ссылка со 2-го слайда на 3-й — после нарезки 3-го слайда не будет
    prs.slides[1].shapes[0].click_action.target_slide = prs.slides[2]
    path = tmp_path / "игра.pptx"
    prs.save(path)
    return path


def test_extract_keeps_only_one_slide_and_its_picture(three_slides, tmp_path):
    out = tmp_path / "песня.pptx"

    extract_slide(str(three_slides), "ppt/slides/slide2.xml", str(out))

    prs = Presentation(str(out))
    assert len(prs.slides) == 1
    assert prs.slides[0].shapes[0].text_frame.text == "Слайд 2"
    with zipfile.ZipFile(out) as z:
        media = [n for n in z.namelist() if n.startswith("ppt/media/")]
        slides = [n for n in z.namelist() if n.startswith("ppt/slides/slide")]
    assert len(media) == 1 and slides == ["ppt/slides/slide2.xml"]


def test_link_to_removed_slide_points_to_itself(three_slides, tmp_path):
    out = tmp_path / "песня.pptx"

    extract_slide(str(three_slides), "ppt/slides/slide2.xml", str(out))

    slide = Presentation(str(out)).slides[0]
    assert slide.shapes[0].click_action.target_slide.part is slide.part


def test_unknown_slide_is_an_error(three_slides, tmp_path):
    with pytest.raises(ValueError, match="slide9"):
        extract_slide(str(three_slides), "ppt/slides/slide9.xml", str(tmp_path / "x.pptx"))


REAL_GAME = f"{paths.OLD_DESIGN}/Тематичекие игры/Хэллоуин/Prezentatsia_3_igra.pptx"
REAL_SLIDE = "ppt/slides/slide20.xml"  # «Звери - Просто такая сильная любовь», звук с обрезкой


@pytest.mark.integration
def test_real_old_slide_keeps_trim_in_powerpoint(tmp_path):
    import win32com.client as win32

    with zipfile.ZipFile(REAL_GAME) as z:
        trim = re.search(r'<p14:trim st="([\d.]+)"', z.read(REAL_SLIDE).decode("utf-8"))
    out = tmp_path / "Звери.pptx"

    extract_slide(REAL_GAME, REAL_SLIDE, str(out))

    assert os.path.getsize(out) < 60_000_000  # не вся игра (≈860 МБ), а один слайд
    app = win32.Dispatch("PowerPoint.Application")
    try:
        prs = app.Presentations.Open(str(out), ReadOnly=True, WithWindow=False)
        try:
            assert prs.Slides.Count == 1
            audio = next(s for s in prs.Slides(1).Shapes if s.Type == 16)
            assert abs(audio.MediaFormat.StartPoint - float(trim.group(1))) < 2
        finally:
            prs.Close()
    finally:
        if app.Presentations.Count == 0:
            app.Quit()
