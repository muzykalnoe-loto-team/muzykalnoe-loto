import json
import os
import shutil

import pytest
import win32com.client as win32
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from scripts import paths
from scripts.old_slot import fill_old_slot, insert_photos, title_text

SLOT_3 = 18  # бочонок 3, макет «Игровой слайд», заготовка «НА 1 СТРОЧКУ»


def _catalog_entry(with_photo=True):
    if not os.path.exists(paths.OLD_CATALOG):
        pytest.skip("банк старых слайдов ещё не собран (Task 3)")
    with open(paths.OLD_CATALOG, encoding="utf-8") as f:
        catalog = json.load(f)
    return next(c for c in catalog if bool(c["photo_part"]) == with_photo and c["has_trim"] and not c["check"])


@pytest.fixture
def template_copy(tmp_path):
    path = tmp_path / "шаблон.pptx"
    shutil.copy(paths.PRESENTATION_TEMPLATE, path)
    return path


def test_insert_photos_fills_picture_placeholder(template_copy):
    entry = _catalog_entry()

    insert_photos(str(template_copy), {SLOT_3: (os.path.join(paths.OLD_BANK, entry["file"]), entry["photo_part"])})

    slide = Presentation(str(template_copy)).slides[SLOT_3 - 1]
    photo = next(p for p in slide.placeholders if p.placeholder_format.type == PP_PLACEHOLDER.PICTURE)
    assert photo.__class__.__name__ == "PlaceholderPicture"  # в рамке картинка, а не пустая заготовка


@pytest.fixture
def opened(template_copy):
    app = win32.Dispatch("PowerPoint.Application")
    prs = app.Presentations.Open(str(template_copy), WithWindow=False)
    try:
        yield app, prs
    finally:
        prs.Close()
        if app.Presentations.Count == 0:
            app.Quit()


def _top_shape_at(slide, x, y):
    hits = [s for s in slide.Shapes if s.Left <= x <= s.Left + s.Width and s.Top <= y <= s.Top + s.Height]
    return max(hits, key=lambda s: s.ZOrderPosition) if hits else None


@pytest.mark.integration
def test_fill_old_slot_sets_texts_audio_and_keeps_corner_free(opened):
    app, prs = opened
    entry = _catalog_entry()
    old_file = os.path.abspath(os.path.join(paths.OLD_BANK, entry["file"]))
    src = app.Presentations.Open(old_file, ReadOnly=True, WithWindow=False)
    start = next(s for s in src.Slides(1).Shapes if s.Type == 16).MediaFormat.StartPoint
    src.Close()

    fill_old_slot(app, prs, SLOT_3, old_file, "Просто такая сильная любовь", "Звери")

    slide = prs.Slides(SLOT_3)
    boxes = sorted((g for s in slide.Shapes if s.Type == 6 for g in s.GroupItems if g.HasTextFrame), key=lambda g: g.Top)
    texts = [b.TextFrame.TextRange.Text for b in boxes]
    assert texts[0] == "ПРОСТО ТАКАЯ СИЛЬНАЯ ЛЮБОВЬ" and texts[1] == "Звери"
    # начертание из заготовок шаблона: название жирное, исполнитель обычный (замечание оператора)
    assert boxes[0].TextFrame.TextRange.Font.Bold and not boxes[1].TextFrame.TextRange.Font.Bold
    audio = next(s for s in slide.Shapes if s.Type == 16)
    assert abs(audio.MediaFormat.StartPoint - start) < 2  # тот же фрагмент
    assert audio.AnimationSettings.PlaySettings.PlayOnEntry
    # как в старых играх: сначала звук, а текст выезжает вместе с ним через паузу (замечание оператора)
    seq = slide.TimeLine.MainSequence
    assert seq(1).Shape.Name == audio.Name
    text_effect = next(seq(i) for i in range(2, seq.Count + 1) if seq(i).Shape.Type == 6)
    assert text_effect.Timing.TriggerType == 2 and text_effect.Timing.TriggerDelayTime > 0  # «с предыдущим», с паузой
    # значок в правом верхнем углу, как в старых играх (замечание оператора)
    assert abs(audio.Left + audio.Width - prs.PageSetup.SlideWidth) < 1 and audio.Top < 20
    assert _top_shape_at(slide, 940, 520) is None  # угол с кнопкой макета свободен


def _boxes(slide):
    return sorted((g for s in slide.Shapes if s.Type == 6 for g in s.GroupItems if g.HasTextFrame), key=lambda g: g.Top)


def _one_line_height(box):
    rng = box.TextFrame.TextRange
    text = rng.Text
    rng.Text = "А"
    line = rng.BoundHeight
    rng.Text = text
    return line


@pytest.mark.integration
def test_two_line_title_keeps_size_and_pushes_artist_down(opened):
    app, prs = opened
    entry = _catalog_entry()
    slide = prs.Slides(SLOT_3)
    title_box, artist_box = _boxes(slide)
    size_before, line = title_box.TextFrame.TextRange.Font.Size, _one_line_height(title_box)

    fill_old_slot(app, prs, SLOT_3, os.path.abspath(os.path.join(paths.OLD_BANK, entry["file"])),
                  "Коламбия Пикчерз не представляет", "БАНД'ЭРОС")

    rng = title_box.TextFrame.TextRange
    assert rng.Font.Size == size_before  # в две строки влезает — шрифт не уменьшаем
    assert line * 1.5 < rng.BoundHeight <= line * 2.2  # ровно две строки
    assert artist_box.Top >= title_box.Top + title_box.Height - 1  # исполнитель не налезает
    assert artist_box.Top - (title_box.Top + title_box.Height) <= 5  # и не отрывается от названия


@pytest.mark.integration
def test_title_longer_than_two_lines_gets_smaller(opened):
    app, prs = opened
    entry = _catalog_entry()
    slide = prs.Slides(SLOT_3)
    title_box, _ = _boxes(slide)
    size_before, line = title_box.TextFrame.TextRange.Font.Size, _one_line_height(title_box)

    fill_old_slot(app, prs, SLOT_3, os.path.abspath(os.path.join(paths.OLD_BANK, entry["file"])),
                  "Очень длинное название песни которое никак не помещается в две строки", "Кто-то")

    rng = title_box.TextFrame.TextRange
    assert rng.Font.Size < size_before
    assert rng.BoundHeight <= line * rng.Font.Size / size_before * 2.2


def test_title_is_upper_and_hyphenated_words_do_not_break():
    # «ПО-» на одной строке и «ФРАНЦУЗСКИ» на другой — некрасиво (замечание на пробном слайде)
    assert title_text("Люби меня по-французски") == "ЛЮБИ МЕНЯ ПО‑ФРАНЦУЗСКИ"
    assert title_text("Ля - ля") == "ЛЯ - ЛЯ"  # тире между словами не трогаем


SLOT_6 = 21  # бочонок 6, заготовка «ТЕКСТ НА 2 СТРОЧКИ»


@pytest.mark.integration
def test_one_line_title_in_two_line_slot_does_not_cover_artist(opened):
    # пробный слайд: «НАШЕ ЛЕТО 2.0» в одну строку наехало на «Баста & МакSим»
    app, prs = opened
    entry = _catalog_entry()
    title_box, artist_box = _boxes(prs.Slides(SLOT_6))
    middle = (title_box.Top + artist_box.Top + artist_box.Height) / 2

    fill_old_slot(app, prs, SLOT_6, os.path.abspath(os.path.join(paths.OLD_BANK, entry["file"])), "Наше лето", "Баста")

    gap = artist_box.Top - (title_box.Top + title_box.Height)
    assert -1 <= gap <= 5  # исполнитель вплотную под названием, не поверх
    assert abs((title_box.Top + artist_box.Top + artist_box.Height) / 2 - middle) < 2  # пара по центру, как в заготовке
