import io
import os

import openpyxl
from PIL import Image
from pptx import Presentation
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.util import Pt

from scripts.old_slides import choose_best, identify_song, load_svod_songs, scan_presentation

AUDIO_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/audio"


def _png(color, size=(40, 30)):
    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, "PNG")
    buf.seek(0)
    return buf


def _write_svod(folder, rows):
    os.makedirs(os.path.join(folder, "Тема"), exist_ok=True)
    wb = openpyxl.Workbook()
    wb.active.title = "СОДЕРЖАНИЕ"
    ws = wb.create_sheet("Игра")
    for i, (artist, title) in enumerate(rows, 1):
        ws.append([i, f"{artist} - {title}", artist + " ", title])
    wb.save(os.path.join(folder, "Тема", "Тема_СВОД песен.xlsx"))


def _make_old_game(path, slides):
    """slides: список (тексты, цвет_фото | None). На каждом слайде — звук и общая картинка оформления."""
    prs = Presentation()
    decor = _png((0, 0, 0), (10, 10)).getvalue()
    for texts, photo in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        for i, t in enumerate(texts):
            slide.shapes.add_textbox(Pt(20), Pt(20 + 40 * i), Pt(300), Pt(30)).text_frame.text = t
        slide.shapes.add_picture(io.BytesIO(decor), Pt(0), Pt(0), Pt(10), Pt(10))
        if photo:
            slide.shapes.add_picture(_png(photo), Pt(50), Pt(50), Pt(400), Pt(300))
        # звук: связь с mp3 — этого достаточно, чтобы слайд считался песенным
        mp3 = Part(PackURI(f"/ppt/media/media{len(prs.slides)}.mp3"), "audio/mpeg", prs.part.package, b"ID3fake")
        slide.part.relate_to(mp3, AUDIO_REL)
    prs.save(path)


def test_identify_song_prefers_longest_match_and_flags_unknown():
    svod = [
        {"raw": "Андрей Губин - Ночь", "artist_key": "андрейгубин", "title_key": "ночь"},
        {"raw": "Юрий Шатунов - Седая ночь", "artist_key": "юрийшатунов", "title_key": "седаяночь"},
    ]
    assert identify_song(["70", "СЕДАЯ НОЧЬ", "Юрий Шатунов"], svod) == ("Юрий Шатунов - Седая ночь", "")
    raw, check = identify_song(["28", "РОЗОВЫЙ ФЛАМИГО", "Алёна Свиридова"], svod)
    assert raw == "Алёна Свиридова - Розовый фламиго"
    assert "проверьте" in check
    assert identify_song(["12"], svod) == (None, "")


def test_scan_finds_song_slides_and_photo(tmp_path):
    _write_svod(str(tmp_path), [("Звери", "Танцуй"), ("Земфира", "Искала")])
    game = tmp_path / "Тема" / "Игра" / "1.pptx"
    os.makedirs(game.parent)
    _make_old_game(str(game), [
        (["1", "ТАНЦУЙ", "Звери"], (255, 0, 0)),
        (["2", "ИСКАЛА", "Земфира"], None),  # без своего фото
        (["3", "ИСКАЛА", "Земфира"], (0, 0, 255)),
    ])
    svod = load_svod_songs(str(tmp_path))

    found = scan_presentation(str(game), svod)

    assert [f["raw"] for f in found] == ["Звери - Танцуй", "Земфира - Искала", "Земфира - Искала"]
    assert found[0]["photo_part"] is not None
    assert found[1]["photo_part"] is None  # общая картинка оформления — не фото
    best = choose_best(found)
    zemfira = next(b for b in best if b["raw"] == "Земфира - Искала")
    assert zemfira["photo_part"] is not None  # из дублей выбран слайд с фото
    assert len(zemfira["spares"]) == 1


def test_photo_inside_layout_placeholder_is_found(tmp_path):
    # В старых играх фото часто вставлено в рамку макета: своего размера у картинки нет
    _write_svod(str(tmp_path), [("Иракли", "Вова-чума"), ("Звери", "Танцуй")])
    game = tmp_path / "Тема" / "Игра" / "1.pptx"
    os.makedirs(game.parent)
    prs = Presentation()
    for i, (texts, color) in enumerate([(["1", "ВОВА-ЧУМА", "Иракли"], (200, 0, 0)), (["2", "ТАНЦУЙ", "Звери"], (0, 200, 0))]):
        slide = prs.slides.add_slide(prs.slide_layouts[8])  # «Рисунок с подписью» — есть рамка под фото
        slide.placeholders[1].insert_picture(_png(color))
        for j, t in enumerate(texts):
            slide.shapes.add_textbox(Pt(20), Pt(20 + 40 * j), Pt(300), Pt(30)).text_frame.text = t
        mp3 = Part(PackURI(f"/ppt/media/media{i}.mp3"), "audio/mpeg", prs.part.package, b"ID3fake")
        slide.part.relate_to(mp3, AUDIO_REL)
    prs.save(game)

    found = scan_presentation(str(game), load_svod_songs(str(tmp_path)))

    assert [f["photo_part"] is not None for f in found] == [True, True]
