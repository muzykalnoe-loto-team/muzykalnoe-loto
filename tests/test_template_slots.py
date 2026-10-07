import pytest
from pptx import Presentation
from scripts.template_slots import find_slide_index_by_barrel_number

TEMPLATE = "Рабочий 2 вар.pptx"


def test_finds_slot_for_barrel_1():
    prs = Presentation(TEMPLATE)

    index = find_slide_index_by_barrel_number(prs, 1)

    assert index == 15  # слайд №16 в 1-индексации PowerPoint


def test_finds_slot_for_barrel_90():
    prs = Presentation(TEMPLATE)

    index = find_slide_index_by_barrel_number(prs, 90)

    assert index == 104  # слайд №105 в 1-индексации PowerPoint


def test_raises_for_unknown_barrel():
    prs = Presentation(TEMPLATE)

    with pytest.raises(LookupError):
        find_slide_index_by_barrel_number(prs, 91)
