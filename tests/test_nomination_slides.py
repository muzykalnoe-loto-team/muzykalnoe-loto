import shutil

import pytest

from scripts import paths
from scripts.nomination_slides import apply_nominations, title_lines

pytestmark = pytest.mark.integration

PP_HYPERLINK = 7
TABLE_BUTTONS = {1: 505, 2: 594, 3: 679}  # кнопки «1», «2», «3» под сеткой «Таблицы» (Left, pt)


@pytest.fixture
def opened(tmp_path):
    import win32com.client as win32
    path = tmp_path / "игра.pptx"
    shutil.copy(paths.PRESENTATION_TEMPLATE, path)
    app = win32.Dispatch("PowerPoint.Application")
    prs = app.Presentations.Open(str(path), WithWindow=False)
    try:
        yield app, prs
    finally:
        prs.Close()
        if app.Presentations.Count == 0:
            app.Quit()


def _slides(prs):
    return [prs.Slides(i) for i in range(1, prs.Slides.Count + 1)]


def _link(shape):
    try:
        action = shape.ActionSettings(1)
    except Exception:  # у групп нет действия
        return None
    if action.Action != PP_HYPERLINK or not action.Hyperlink.SubAddress:
        return None
    return action.Hyperlink.SubAddress


def link_target(shape):
    sub = _link(shape)
    return int(sub.split(",")[0]) if sub else None


def _table(prs):
    return next(s for s in _slides(prs) if s.CustomLayout.Name == "Таблица")


def button_targets(prs):
    table, found = _table(prs), {}
    for shape in table.Shapes:
        for n, left in TABLE_BUTTONS.items():
            if abs(shape.Left - left) < 5 and shape.Top > 450:
                found[n] = link_target(shape)
    return found


def test_regular_game_keeps_one_per_round_and_links_buttons(opened):
    app, prs = opened
    apply_nominations(prs, {1: "cry", 2: "cherry", 3: "dancefloor"}, None)
    ids = {s.SlideID for s in _slides(prs)}
    assert {258, 376, 377} <= ids and not ({374, 375, 259, 378} & ids)
    assert prs.Slides.Count == 118 - 4
    assert button_targets(prs) == {1: 258, 2: 376, 3: 377}


def test_no_link_leads_to_a_deleted_or_shifted_slide(opened):
    # кнопки «дальше» на номинациях и «Розыгрыше» в шаблоне ведут на конкретные варианты
    # (например, «Крик души» → «Пьяная вишня»): удалённый вариант заменяется оставшимся в том же раунде
    app, prs = opened
    apply_nominations(prs, {1: "cry", 2: "goddess", 3: "dancefloor"}, None)
    index_of = {s.SlideID: s.SlideIndex for s in _slides(prs)}
    layouts = {s.CustomLayout.Name: s.CustomLayout for s in _slides(prs)}.values()
    for owner in _slides(prs) + list(layouts):
        for shape in owner.Shapes:
            sub = _link(shape)
            if sub:
                slide_id, index = (int(x) for x in sub.split(",")[:2])
                assert index_of.get(slide_id) == index, f"{sub} на «{owner.Name}»"
    assert link_target(_corner_button(prs.Slides.FindBySlideID(258))) == 259  # «Крик души» → «Синяя богиня»


def _corner_button(slide):
    return next(s for s in slide.Shapes if link_target(s) and s.Left > 850)


def test_slot_buttons_still_lead_to_table(opened):
    app, prs = opened
    apply_nominations(prs, {1: "sing", 2: "goddess", 3: "nobrakes"}, None)
    slot = next(s for s in _slides(prs) if s.CustomLayout.Name.startswith("Игровой слайд"))  # бочонок 1
    assert any(link_target(s) == _table(prs).SlideID for s in slot.Shapes)


def test_themed_game_fills_slide_and_drops_regular_round1(opened, tmp_path):
    from scripts.theme_art import render_scene
    from scripts.theme_scenes import THEMES
    app, prs = opened
    gatsby = THEMES["gatsby"]
    art = render_scene(gatsby["scene"], str(tmp_path / "g.png"), seed="gatsby")
    apply_nominations(prs, {1: "theme", 2: "goddess", 3: "dancefloor"},
                      {"title": gatsby["title"], "glow": gatsby["glow"], "art_png": art})
    slide = prs.Slides.FindBySlideID(374)
    texts = [s.TextFrame.TextRange.Text for s in slide.Shapes if s.HasTextFrame and s.TextFrame.HasText]
    assert any("ГЭТСБИ" in t for t in texts) and any("САМАЯ КОСТЮМИРОВАННАЯ КОМАНДА" in t for t in texts)
    assert not any("РОК" in t for t in texts)
    assert title_lines(slide) == 2
    assert button_targets(prs)[1] == 374
    full = [s for s in slide.Shapes if s.Type == 13 and s.Width > 900]
    assert len(full) == 1 and full[0].ZOrderPosition == 1  # одна картинка на весь слайд, под всем остальным
    assert _corner_button(slide).ZOrderPosition > 1  # кнопка в углу поверх картинки


def _variants():
    from scripts.theme_scenes import THEMES
    return [{"title": "БЛЕСК\vПО‑ГЭТСБИ", "glow": "#E8B64C", "scene": THEMES["gatsby"]["scene"], "seed": "a"},
            {"title": "ТЬМА ВАМ\vК ЛИЦУ", "glow": "#8A3DFF", "scene": THEMES["gothic"]["scene"], "seed": "b"}]


def test_preview_has_two_slides_and_pictures(tmp_path):
    import os

    from pptx import Presentation

    from scripts.nomination_slides import preview_variants
    res = preview_variants(_variants(), str(tmp_path))
    assert len(res["pngs"]) == 2 and all(os.path.exists(p) for p in res["pngs"] + res["arts"])
    assert len(Presentation(res["pptx"]).slides) == 2 and res["warnings"] == []


def test_too_long_title_is_reported(tmp_path):
    from scripts.nomination_slides import preview_variants
    v = [dict(_variants()[0], title="ОЧЕНЬ ДЛИННЫЙ ЗАГОЛОВОК\vКОТОРЫЙ НЕ ВЛЕЗЕТ")]
    warnings = preview_variants(v, str(tmp_path))["warnings"]
    assert len(warnings) == 1 and "строк" in warnings[0] and "1" in warnings[0]


def test_archive_grows_by_one_slide_with_note(tmp_path):
    # в тесте архив читать можно — запрет касается настоящего архива
    from pptx import Presentation

    from scripts.nomination_slides import append_to_archive, preview_variants
    game = preview_variants(_variants()[:1], str(tmp_path))["pptx"]  # в нём тематический слайд (SlideID 374)
    archive = str(tmp_path / "Номинация.pptx")
    assert append_to_archive(game, "Гэтсби-бал", "2026-11-01", "Гэтсби", archive) == ""
    assert append_to_archive(game, "Ещё бал", "2026-12-01", "Гэтсби", archive) == ""
    prs = Presentation(archive)
    assert len(prs.slides) == 2
    assert "Ещё бал — 2026-12-01 — Гэтсби" in prs.slides[-1].notes_slide.notes_text_frame.text
    assert any("ГЭТСБИ" in s.text_frame.text for s in prs.slides[-1].shapes if s.has_text_frame)


def test_busy_archive_gives_warning_not_error(tmp_path, opened):
    # архив открыт в PowerPoint (у оператора) — игра собирается, дописать нельзя: предупреждение
    from scripts.nomination_slides import append_to_archive, preview_variants
    app, prs = opened
    game = preview_variants(_variants()[:1], str(tmp_path))["pptx"]
    warning = append_to_archive(game, "Бал", "2026-11-01", "Гэтсби", prs.FullName)
    assert "архив" in warning.lower()
