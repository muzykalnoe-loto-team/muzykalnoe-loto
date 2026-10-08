# Общий банк песен: видео-караоке (новый банк) + старые слайды (фото + звук) в одном списке.
# Запуск из корня проекта: python -m scripts.common_bank
import json
import os

import openpyxl
from openpyxl.styles import Font, PatternFill

from scripts import paths
from scripts.song_bank import load_song_bank
from scripts.song_names import norm


def _blank(raw: str) -> dict:
    return {"raw": raw, "video_file": None, "old_file": None, "old_photo_part": None, "old_source": "", "check": ""}


def merge_banks(video_songs: list[str], video_files: set[str], old_catalog: list[dict]) -> list[dict]:
    entries: dict[str, dict] = {}
    for raw in video_songs:
        entry = entries.setdefault(norm(raw), _blank(raw))
        if f"{raw}.pptx" in video_files:
            entry["video_file"] = f"{raw}.pptx"
    for old in old_catalog:
        if not old["file"]:  # слайд не нарезан (ждёт проверки) — в игру его взять нельзя
            continue
        entry = entries.setdefault(norm(old["raw"]), _blank(old["raw"]))
        entry.update(old_file=old["file"], old_photo_part=old["photo_part"],
                     old_source=f"{old['source'].replace(paths.OLD_DESIGN + '/', '')} / {old['slide_part']}",
                     check=old["check"])
    return sorted(entries.values(), key=lambda e: e["raw"].lower())


def build_common_bank(song_bank: str = paths.SONG_BANK, archive: str = paths.SLIDES_ARCHIVE,
                      old_catalog: str = paths.OLD_CATALOG) -> list[dict]:
    with open(old_catalog, encoding="utf-8") as f:
        old = json.load(f)
    video = [s["raw"] for s in load_song_bank(song_bank)]
    return merge_banks(video, set(os.listdir(archive)), old)


def save_common_bank(entries: list[dict], xlsx_path: str = paths.COMMON_BANK_XLSX,
                     json_path: str = paths.COMMON_BANK_JSON) -> None:
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(entries, f, ensure_ascii=False, indent=1)
    wb = openpyxl.Workbook()
    header = ["Исполнитель - Название", "Видео (караоке)", "Старый слайд (фото + звук)", "Откуда старый слайд", "Проверить"]
    rows = [[e["raw"], "да" if e["video_file"] else "нет", "да" if e["old_file"] else "нет", e["old_source"], e["check"]]
            for e in entries]
    for i, (title, data) in enumerate((("Все песни", rows), ("Проверить", [r for r in rows if r[4]]))):
        sheet = wb.active if i == 0 else wb.create_sheet()
        sheet.title = title
        sheet.append(header)
        for r in data:
            sheet.append(r)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="FFE699")
        for col, width in zip("ABCDE", (50, 16, 26, 80, 40)):
            sheet.column_dimensions[col].width = width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    wb.save(xlsx_path)


def load_common_bank(json_path: str = paths.COMMON_BANK_JSON) -> list[dict]:
    with open(json_path, encoding="utf-8") as f:
        return json.load(f)


if __name__ == "__main__":
    bank = build_common_bank()
    save_common_bank(bank)
    both = sum(1 for e in bank if e["video_file"] and e["old_file"])
    print(f"песен: {len(bank)}; с видео: {sum(1 for e in bank if e['video_file'])}; "
          f"со старым слайдом: {sum(1 for e in bank if e['old_file'])}; в обоих видах: {both}")
