# Банк старых слайдов: каждая песня из старых игр — отдельный файл, плюс каталог и отчёт.
# Запуск из корня проекта: python -m scripts.build_old_bank
import glob
import json
import os
import re
import zipfile

import openpyxl
from openpyxl.styles import Font, PatternFill

from scripts import paths
from scripts.compact_slide import compact_slide
from scripts.old_slides import choose_best, load_svod_songs, scan_presentation
from scripts.slide_split import extract_slide
from scripts.song_names import safe_filename

REPORT = f"{paths.SOURCES}/Банк старых слайдов — отчёт.xlsx"
WAITING = "ждёт проверки"
# Правки оператора по песням «проверить»: «источник|слайд» или «raw|написание» → новое название или DROP
FIXES = f"{paths.OLD_BANK}/правки оператора.json"
DROP = "удалить"


def apply_fixes(entries: list[dict], fixes: dict[str, str]) -> list[dict]:
    result = []
    for entry in entries:
        # правка конкретного слайда или любого слайда с тем же (ошибочным) написанием
        fix = fixes.get(f"{entry['source']}|{entry['slide_part']}") or fixes.get(f"raw|{entry['raw']}")
        if fix == DROP:
            continue
        if fix:
            entry = dict(entry, raw=fix, check="")
        result.append(entry)
    return result


def load_fixes(path: str = FIXES) -> dict[str, str]:
    if not os.path.exists(path):
        return {}
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def assign_file_names(catalog: list[dict]) -> None:
    used = set()
    for entry in catalog:
        name = safe_filename(entry["raw"])
        while name.lower() in used:
            name += "_"
        used.add(name.lower())
        entry["file"] = name + ".pptx"


def songs_to_cut(catalog: list[dict]) -> list[dict]:
    # Песни с пометкой «проверить» режем только после сверки оператором (решение 2026-10-08)
    return [c for c in catalog if not c["check"]]


def is_cut_well(path: str) -> bool:
    """Файл банка в порядке: открывается и в нём ровно один слайд."""
    try:
        with zipfile.ZipFile(path) as z:
            return sum(1 for n in z.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)) == 1
    except (OSError, zipfile.BadZipFile):
        return False


def _has_part(path: str, part: str) -> bool:
    try:
        with zipfile.ZipFile(path) as z:
            return part in z.namelist()
    except (OSError, zipfile.BadZipFile):
        return False


def _load_catalog(path: str) -> list[dict]:
    if not os.path.exists(path):
        return []
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def prepare_bank_file(entry: dict, bank_dir: str, previous: dict[str, dict]) -> None:
    """Файл песни в банке: нарезан из старой игры и сжат (звук — отрывок с запасом, фото — JPEG).
    Повторный запуск режет только недостающее и испорченное и не сжимает уже сжатое."""
    out = os.path.join(bank_dir, entry["file"])
    before = previous.get(entry["file"], {})
    cut_from = f"{entry['source']}|{entry['slide_part']}"
    entry["cut_from"] = cut_from
    # файл нарезан из того же слайда, что выбран сейчас (лучший слайд мог смениться после правок узнавания);
    # для файлов, нарезанных до записи «cut_from», сверяем номер слайда внутри файла
    same_slide = before["cut_from"] == cut_from if before.get("cut_from") else _has_part(out, entry["slide_part"])
    cut_well = is_cut_well(out) and same_slide
    if cut_well and before.get("compact"):
        entry.update(compact=True, photo_part=before["photo_part"])
        return
    tmp = out + ".part"
    if not cut_well:
        extract_slide(entry["source"], entry["slide_part"], tmp)
        os.replace(tmp, out)  # недорезанный файл не выдаст себя за готовый
    try:
        new_photo = compact_slide(out, tmp, entry["photo_part"])
        os.replace(tmp, out)
        entry.update(compact=True, photo_part=new_photo)
    except Exception as error:  # странный файл — оставляем несжатым, песня в банке остаётся
        if os.path.exists(tmp):
            os.remove(tmp)
        entry.update(compact=False, compact_error=f"{type(error).__name__}: {error}")


def _write_report(catalog: list[dict], path: str) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Сводка"
    ws.append(["Песен в каталоге", len(catalog)])
    ws.append(["Нарезано", sum(1 for c in catalog if c["file"])])
    ws.append(["Ждут проверки (не нарезаны)", sum(1 for c in catalog if not c["file"])])
    ws.append(["Без фото", sum(1 for c in catalog if not c["photo_part"])])
    ws.append(["Сжато", sum(1 for c in catalog if c.get("compact"))])
    ws.append(["Не удалось сжать (лежат несжатыми)", sum(1 for c in catalog if c.get("compact_error"))])
    ws.column_dimensions["A"].width = 32
    for title, rows in (("Все песни", catalog), ("Проверить", [c for c in catalog if c["check"]])):
        sheet = wb.create_sheet(title)
        sheet.append(["Песня", "Файл", "Проверить", "Фото", "Источник", "Слайд", "Запасных"])
        for c in rows:
            sheet.append([c["raw"], c["file"] or WAITING, c["check"], "да" if c["photo_part"] else "нет",
                          c["source"].replace(paths.OLD_DESIGN + "/", ""), c["slide_part"], len(c["spares"])])
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="FFE699")
        for col, width in zip("ABCDEFG", (48, 50, 40, 8, 70, 24, 10)):
            sheet.column_dimensions[col].width = width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    wb.save(path)


def build_old_bank(old_design: str = paths.OLD_DESIGN, bank_dir: str = paths.OLD_BANK,
                   catalog_path: str = paths.OLD_CATALOG, log=print) -> list[dict]:
    svod = load_svod_songs(old_design)
    entries = []
    for pres in sorted(glob.glob(os.path.join(old_design, "*", "*", "*.pptx"))):
        entries += scan_presentation(pres, svod)
    catalog = choose_best(apply_fixes(entries, load_fixes()))
    assign_file_names(catalog)
    to_cut = songs_to_cut(catalog)
    cut_ids = {id(c) for c in to_cut}
    for entry in catalog:
        if id(entry) not in cut_ids:
            entry["file"] = None  # в банке файла нет — песня ждёт проверки
    os.makedirs(bank_dir, exist_ok=True)
    previous = {c["file"]: c for c in _load_catalog(catalog_path) if c.get("file")}
    for i, entry in enumerate(to_cut, 1):
        prepare_bank_file(entry, bank_dir, previous)
        if i % 100 == 0:
            log(f"готово {i} из {len(to_cut)}")
    with open(catalog_path, "w", encoding="utf-8") as f:
        json.dump(catalog, f, ensure_ascii=False, indent=1)
    _write_report(catalog, REPORT)
    return catalog


if __name__ == "__main__":
    result = build_old_bank(log=lambda m: print(m, flush=True))
    print(f"Готово: нарезано {sum(1 for c in result if c['file'])} песен, ждут проверки {sum(1 for c in result if not c['file'])}, "
          f"сжато {sum(1 for c in result if c.get('compact'))}, не удалось сжать {sum(1 for c in result if c.get('compact_error'))}")
