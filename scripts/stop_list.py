# Стоп-лист: исполнители, которых не берём в игры (например, иноагенты). Список ведёт оператор
# в Excel (`paths.STOP_LIST`, первая колонка; в одной ячейке можно несколько написаний через «/»)
import os
import re

from openpyxl import load_workbook

from scripts import paths
from scripts.song_names import norm, split_raw

# соавторы в строке исполнителя: «А, Б», «А feat. Б», «А & Б», «А и Б», «А x Б»
_SEPARATORS = re.compile(r",|\bfeat\.?|\bft\.?|&|\s+и\s+|\s+x\s+", re.IGNORECASE)


def load_stop_list(path: str | None = None) -> list[str]:
    path = path or paths.STOP_LIST
    if not os.path.exists(path):
        return []
    ws = load_workbook(path, read_only=True).worksheets[0]
    names = [str(row[0]).strip() for row in ws.iter_rows(min_row=2, max_col=1, values_only=True) if row[0]]
    return [n for n in names if n]


def blocked_artist(raw: str, stop_list: list[str]) -> str | None:
    """Строка стоп-листа, под которую попадает исполнитель песни (целым именем, в том числе в дуэте)."""
    artists = {norm(a) for a in _SEPARATORS.split(split_raw(raw)[0]) if norm(a)}
    for entry in stop_list:
        if any(norm(name) in artists for name in entry.split("/") if norm(name)):
            return entry
    return None
