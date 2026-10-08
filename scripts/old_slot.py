# Слот «Игровой слайд» шаблона «Рабочий 2 вар» ← песня из банка старых слайдов:
# фото в рамку (python-pptx), название/исполнитель и звук с автозапуском (PowerPoint COM)
import io
import os
import re
import zipfile

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from scripts.ppt_paste import copy_paste_checked

MSO_GROUP = 6
MSO_MEDIA = 16
MAX_TITLE_LINES = 2  # правило оператора: название — не больше двух строк, иначе меньше шрифт
MIN_TITLE_SIZE = 20
MIN_ARTIST_SIZE = 16
LINE_SLACK = 1.1  # запас на неровную высоту строк


def insert_photos(pptx_path: str, photos: dict[int, tuple[str, str]]) -> None:
    if not photos:
        return
    prs = Presentation(pptx_path)
    for slide_number, (old_file, photo_part) in photos.items():
        slide = prs.slides[slide_number - 1]
        frame = next((p for p in slide.placeholders if p.placeholder_format.type == PP_PLACEHOLDER.PICTURE), None)
        if frame is None:
            raise LookupError(f"На слайде {slide_number} нет рамки под фото — это не слот «Игровой слайд»")
        with zipfile.ZipFile(old_file) as z:
            frame.insert_picture(io.BytesIO(z.read(photo_part)))  # обрезка по рамке без искажений
    prs.save(pptx_path)


def title_text(title: str) -> str:
    # дефис внутри слова — неразрывный: «ПО-ФРАНЦУЗСКИ» переносится целиком
    return re.sub(r"(?<=\w)-(?=\w)", "\u2011", title.upper())


def _text_boxes(slide):
    for shape in slide.Shapes:
        if shape.Type == MSO_GROUP:
            boxes = [g for g in shape.GroupItems if g.HasTextFrame]
            if len(boxes) == 2:
                return sorted(boxes, key=lambda b: b.Top)  # сверху название, снизу исполнитель
    raise LookupError("В слоте нет заготовки названия и исполнителя")


def _fit_lines(box, text: str, max_lines: int, min_size: float) -> None:
    """Вписывает текст не больше чем в max_lines строк, уменьшая шрифт только если не влезает."""
    rng = box.TextFrame.TextRange
    size0 = rng.Font.Size
    rng.Text = "А"
    line0 = rng.BoundHeight  # высота одной строки при исходном размере
    rng.Text = text
    while rng.BoundHeight > max_lines * line0 * rng.Font.Size / size0 * LINE_SLACK and rng.Font.Size > min_size:
        rng.Font.Size = rng.Font.Size - 2


def _put_title_and_artist(slide, title: str, artist: str) -> None:
    title_box, artist_box = _text_boxes(slide)
    middle = (title_box.Top + artist_box.Top + artist_box.Height) / 2
    gap = artist_box.Top - (title_box.Top + title_box.Height)
    _fit_lines(title_box, title_text(title), MAX_TITLE_LINES, MIN_TITLE_SIZE)
    _fit_lines(artist_box, artist, 1, MIN_ARTIST_SIZE)
    # Рамки сами подстраиваются под текст (одна строка в «2 СТРОЧКИ» или две в «НА 1 СТРОЧКУ»):
    # ставим пару «название + исполнитель» заново — с прежним зазором и по прежнему центру
    title_box.Top = middle - (title_box.Height + gap + artist_box.Height) / 2
    artist_box.Top = title_box.Top + title_box.Height + gap


def fill_old_slot(ppt_app, prs, slide_number: int, old_file: str, title: str, artist: str) -> None:
    slide = prs.Slides(slide_number)
    _put_title_and_artist(slide, title, artist)

    old = ppt_app.Presentations.Open(os.path.abspath(old_file), ReadOnly=True, WithWindow=False)
    try:
        audio = next((s for s in old.Slides(1).Shapes if s.Type == MSO_MEDIA), None)
        if audio is None:
            raise LookupError(f"В «{os.path.basename(old_file)}» нет звука")
        # обрезка, затухание и громкость хранятся в самой фигуре звука и переносятся вместе с ней
        pasted = copy_paste_checked(audio.Copy, slide, audio.Name, MSO_MEDIA)
    finally:
        old.Close()

    sound = pasted(1)
    sound.Left = -sound.Width - 10  # значок за левым краем: не виден и не закрывает кнопку макета
    sound.Top = 0
    play = sound.AnimationSettings.PlaySettings
    play.PlayOnEntry = True  # звук сам включается при открытии слайда
    play.HideWhileNotPlaying = True
